from pathlib import Path
import json

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT / '源代码模板.docx'
FINAL = ROOT / '洛书量子编程语言运行平台-源代码节选-60页.docx'
MANIFEST = ROOT / '.source-docx-work' / 'source-manifest.json'
MAX_LINE_CHARS = 80


def protected(path):
    name = str(path).replace('\\', '/').lower()
    return 'node_modules/' in name or 'quest' in name


def remove_inline_block_comment(line):
    while '/*' in line:
        start = line.find('/*')
        end = line.find('*/', start + 2)
        if end < 0:
            return line[:start].rstrip()
        line = (line[:start] + line[end + 2:]).rstrip()
    return line


def clean_file(path, kind):
    lines = []
    in_block = False
    for raw in Path(path).read_text(encoding='utf-8', errors='ignore').splitlines():
        line = raw.rstrip()
        stripped = line.strip()
        if '/*' in line and '*/' in line:
            line = remove_inline_block_comment(line)
            stripped = line.strip()
        if in_block:
            if '*/' in stripped:
                in_block = False
            continue
        if stripped.startswith('/*'):
            if '*/' not in stripped[2:]:
                in_block = True
            continue
        if kind in {'js', 'c', 'losu'} and stripped.startswith('//'):
            continue
        if kind == 'sh' and stripped.startswith('#'):
            continue
        if kind == 'html' and stripped.startswith('<!--'):
            if '-->' not in stripped[4:]:
                in_block = True
            continue
        if not stripped:
            continue
        if kind == 'sh' and '  #' in line:
            line = line.split('  #', 1)[0].rstrip()
        if line.strip() and len(line) <= MAX_LINE_CHARS:
            lines.append(line)
    return lines


def records_for_files(files, kind):
    records = []
    for path in files:
        if protected(path):
            continue
        rel = path.relative_to(ROOT).as_posix()
        records.extend((line, rel) for line in clean_file(path, kind))
    return records


def inline_records(path, tag, kind):
    text = Path(path).read_text(encoding='utf-8', errors='ignore')
    if tag == 'script':
        body = text.split('<script>', 1)[1].split('</script>', 1)[0]
    elif tag == 'style':
        body = text.split('<style>', 1)[1].split('</style>', 1)[0]
    else:
        body = text
    rel = Path(path).relative_to(ROOT).as_posix() + f'::{tag}'
    return [(line, rel) for line in clean_file_from_text(body, kind) if len(line) <= MAX_LINE_CHARS]


def wasm_binary_records(path):
    data = Path(path).read_bytes()
    rel = Path(path).relative_to(ROOT).as_posix() + '::hex'
    records = []
    for offset in range(0, len(data), 16):
        block = data[offset:offset + 16]
        hex_bytes = ' '.join(f'{value:02x}' for value in block)
        records.append((f'{offset:08x}: {hex_bytes}', rel))
    return records


def clean_file_from_text(text, kind):
    lines = []
    in_block = False
    for raw in text.splitlines():
        line = raw.rstrip()
        stripped = line.strip()
        if '/*' in line and '*/' in line:
            line = remove_inline_block_comment(line)
            stripped = line.strip()
        if in_block:
            if (kind == 'html' and '-->' in stripped) or (kind != 'html' and '*/' in stripped):
                in_block = False
            continue
        if kind == 'html' and stripped.startswith('<!--'):
            if '-->' not in stripped[4:]:
                in_block = True
            continue
        if kind != 'html' and stripped.startswith('/*'):
            if '*/' not in stripped[2:]:
                in_block = True
            continue
        if kind in {'js', 'c', 'losu'} and stripped.startswith('//'):
            continue
        if not stripped:
            continue
        if kind == 'sh' and '  #' in line:
            line = line.split('  #', 1)[0].rstrip()
        if line.strip():
            lines.append(line)
    return lines


def category_records():
    www = ROOT / 'www/index.html'
    js_files = [ROOT / 'index.js', ROOT / 'test.js',
                ROOT / 'src/QLosu/Qapp/utils/docs/layout/doxygen-awesome-darkmode-toggle.js']
    js = records_for_files(js_files, 'js') + inline_records(www, 'script', 'js')

    shell_files = [ROOT / 'src/wasm-build/build.sh', ROOT / 'src/QLosu/Qapp/utils/scripts/compile.sh']
    losu_files = sorted(p for p in ROOT.rglob('*.losu') if not protected(p))
    wasm = records_for_files(shell_files, 'sh')
    wasm += records_for_files(losu_files, 'losu')
    wasm += records_for_files([ROOT / 'src/QLosu/Qapp/src/script/src/losu_config.h'], 'c')
    wasm = wasm_binary_records(ROOT / 'wasm/qlosu.m.wasm') + wasm

    c_files = sorted(p for p in ROOT.rglob('*') if p.is_file() and p.suffix in {'.c', '.h'} and not protected(p))
    c = records_for_files(c_files, 'c')

    html_files = [ROOT / 'www/index.html', ROOT / 'src/QLosu/Qapp/utils/docs/layout/header.html']
    html = records_for_files(html_files, 'html')

    css_files = sorted(p for p in ROOT.rglob('*.css') if not protected(p))
    css = records_for_files(css_files, 'css')
    css += inline_records(www, 'style', 'css')
    categories = {'JavaScript': js, 'WASM/Losu': wasm, 'C': c, 'HTML': html, 'CSS': css}
    return {
        kind: [(line, source) for line, source in lines if 'quest' not in line.lower()]
        for kind, lines in categories.items()
    }


def build_pages():
    quotas = [('JavaScript', 6), ('WASM/Losu', 36), ('C', 8), ('HTML', 5), ('CSS', 5)]
    records = category_records()
    pages = []
    for kind, page_count in quotas:
        needed = page_count * 50
        if len(records[kind]) < needed:
            raise ValueError(f'{kind} only has {len(records[kind])} usable lines; need {needed}')
        for offset in range(0, needed, 50):
            chunk = records[kind][offset:offset + 50]
            pages.append((kind, chunk))
    return pages


def validate_pages(pages):
    assert len(pages) == 60
    for kind, page in pages:
        assert len(page) == 50
        assert all(line.strip() for line, _ in page)
        for line, _ in page:
            assert len(line) <= MAX_LINE_CHARS
            assert not any(marker in line for marker in ('/*', '*/', '<!--', '-->'))
            assert not line.lstrip().startswith('//')


def set_simsun_12(run):
    run.font.name = '宋体'
    run.font.size = Pt(12)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:eastAsia'), '宋体')
    fonts.set(qn('w:ascii'), '宋体')
    fonts.set(qn('w:hAnsi'), '宋体')


def clear_body(doc):
    body = doc._element.body
    sect = body.sectPr
    for child in list(body):
        if child is not sect:
            body.remove(child)


def update_header(doc):
    paragraph = doc.sections[0].header.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = '洛书量子编程语言运行平台 V1.0.0'
        for run in paragraph.runs[1:]:
            run.text = ''
    else:
        paragraph.add_run('洛书量子编程语言运行平台 V1.0.0')


def build():
    pages = build_pages()
    validate_pages(pages)
    doc = Document(str(REFERENCE))
    clear_body(doc)
    update_header(doc)
    style = doc.styles['HTML Preformatted']
    style.font.name = '宋体'
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    style._element.rPr.rFonts.set(qn('w:ascii'), '宋体')
    style._element.rPr.rFonts.set(qn('w:hAnsi'), '宋体')
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = Pt(14.6)

    manifest_pages = []
    for page_index, (kind, records) in enumerate(pages):
        sources = []
        for line_index, (line, source) in enumerate(records):
            paragraph = doc.add_paragraph(style='HTML Preformatted')
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = Pt(14.6)
            paragraph.paragraph_format.keep_together = True
            if line_index == 0 and page_index > 0:
                paragraph.paragraph_format.page_break_before = True
            run = paragraph.add_run(line)
            set_simsun_12(run)
            if source not in sources:
                sources.append(source)
        manifest_pages.append({'page': page_index + 1, 'type': kind, 'sources': sources, 'lines': 50})

    doc.core_properties.title = '洛书量子编程语言运行平台源代码节选（60页）'
    doc.core_properties.subject = 'JavaScript、WebAssembly、C、HTML、CSS 源代码节选'
    doc.core_properties.author = 'QLosu 项目组'
    doc.save(str(FINAL))
    manifest = {
        'output': str(FINAL),
        'page_count': 60,
        'lines_per_page': 50,
        'pages': manifest_pages,
        'excluded_globs': ['src/**/QuEST/**', 'src/**/quest/**', 'node_modules/**'],
        'binary_encoding': 'wasm/qlosu.m.wasm represented as 16-byte hexadecimal rows',
    }
    MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    print(FINAL)


if __name__ == '__main__':
    build()
